"""/app/exportar/ — dentist downloads his own data (CSV with BOM and ';', Excel, clinic backup ZIP)."""

import csv
import io

from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.translation import gettext as _
from django.views.decorators.http import require_POST

from apps.agenda.models import Appointment
from apps.core.activity import log_action
from apps.core.permissions import dentist_required

from .models import Invoice, PatientProfile


def _d(v):
    if not v:
        return ""
    if hasattr(v, "hour"):
        return f"{timezone.localtime(v):%d/%m/%Y %H:%M}"
    return f"{v:%d/%m/%Y}"


def _money(v):
    if v in (None, ""):
        return ""
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _yn(v):
    return _("sim") if v else _("não")


def patients_rows(dentist, full_cpf=False):
    from django.db.models import Sum

    from apps.core.templatetags.wb import mask_cpf

    now = timezone.now()
    header = [_("Nome"), _("E-mail"), _("Telefone"), _("WhatsApp"), _("Data de nascimento"), _("CPF"), _("Entrou no portal"),
              _("Convidado em"), _("Primeira consulta"), _("Última consulta"), _("Próxima consulta"), _("Tratamentos ativos"),
              _("Faturas em aberto (R$)")]
    rows = []
    for p in PatientProfile.objects.for_dentist(dentist).select_related("user"):
        ap = p.appointments.exclude(status__in=["RESCHEDULED", "REJECTED", "EXPIRED", "CANCELLED"])
        first = ap.order_by("start").first()
        last = ap.filter(start__lt=now).order_by("-start").first()
        nxt = ap.filter(start__gte=now).order_by("start").first()
        open_total = Invoice.objects.filter(treatment__patient=p, status="OPEN").aggregate(s=Sum("amount"))["s"]
        rows.append([
            p.name, p.email, p.phone, p.whatsapp, _d(p.birth_date), p.cpf if full_cpf else mask_cpf(p.cpf), _yn(p.has_joined),
            _d(p.invited_at), _d(first.start if first else None), _d(last.start if last else None), _d(nxt.start if nxt else None),
            p.treatments.filter(status__in=["PLANNED", "IN_PROGRESS", "ON_HOLD"]).count(), open_total or 0,
        ])
    return header, rows


def appointments_rows(dentist, start=None, end=None, statuses=None):
    qs = Appointment.objects.for_dentist(dentist).select_related("type", "patient__user", "lead", "treatment").order_by("start")
    if start:
        qs = qs.filter(start__date__gte=start)
    if end:
        qs = qs.filter(start__date__lte=end)
    if statuses:
        qs = qs.filter(status__in=statuses)
    header = [_("Data"), _("Início"), _("Fim"), _("Tipo"), _("Paciente / lead"), _("E-mail"), _("Situação"), _("Origem"),
              _("Tratamento"), _("Observações")]
    rows = []
    for a in qs:
        s, e = timezone.localtime(a.start, dentist.tzinfo), timezone.localtime(a.end, dentist.tzinfo)
        rows.append([f"{s:%d/%m/%Y}", f"{s:%H:%M}", f"{e:%H:%M}", a.type.name, a.person_name, a.person_email, a.get_status_display(),
                     a.get_source_display(), a.treatment.title if a.treatment else "", a.patient_note])
    return header, rows


def leads_rows(dentist):
    header = [_("Data"), _("Nome"), _("E-mail"), _("WhatsApp"), _("Origem"), _("Situação"), _("Mensagem"), _("Anotações")]
    rows = [[_d(x.created_at), x.name, x.email, x.whatsapp, x.get_source_display(), x.get_status_display(), x.message, x.notes]
            for x in dentist.leads.all()]
    return header, rows


def invoices_rows(dentist):
    header = [_("Número"), _("Paciente"), _("Tratamento"), _("Descrição"), _("Valor (R$)"), _("Emissão"), _("Vencimento"),
              _("Situação"), _("Pago em")]
    rows = [[i.display_number, i.treatment.patient.name, i.treatment.title, i.title, i.amount, _d(i.issued_date), _d(i.due_date),
             i.get_status_display(), _d(i.paid_at)]
            for i in Invoice.objects.for_dentist(dentist).select_related("treatment__patient__user")]
    return header, rows


def quotes_rows(dentist):
    header = [_("Título"), _("Destinatário"), _("Total (R$)"), _("Situação"), _("Enviado em"), _("Visitas"),
              _("Tempo de leitura (s)"), _("Respondido em")]
    rows = []
    for q in dentist.quotes.prefetch_related("items", "visits"):
        visits = list(q.visits.all())
        rows.append([q.title, q.recipient_name, q.total, q.get_status_display(), _d(q.sent_at), len(visits),
                     sum(v.duration_s for v in visits), _d(q.responded_at)])
    return header, rows


def _csv(name, header, rows):
    buf = io.StringIO()
    buf.write("﻿")
    w = csv.writer(buf, delimiter=";")
    w.writerow(header)
    for r in rows:
        w.writerow([_money(v) if hasattr(v, "quantize") else v for v in r])
    resp = HttpResponse(buf.getvalue(), content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{name}.csv"'
    return resp


def _xlsx(name, header, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = name[:30]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([float(v) if hasattr(v, "quantize") else v for v in r])
    for i, h in enumerate(header, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:200]]) + 2
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width, 50)
        if "R$" in str(h):
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in cell:
                    c.number_format = '"R$" #,##0.00'
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    return resp


BUILDERS = {
    "pacientes": patients_rows, "consultas": appointments_rows, "leads": leads_rows, "faturas": invoices_rows,
    "orcamentos": quotes_rows,
}


@dentist_required
def exports(request):
    return render(request, "clinic/exports.html")


@dentist_required
def export(request, what, fmt):
    d = request.dentist
    if what not in BUILDERS or fmt not in ("csv", "xlsx"):
        from django.http import Http404

        raise Http404
    if what == "pacientes":
        header, rows = patients_rows(d, full_cpf=bool(request.GET.get("cpf")))
    else:
        header, rows = BUILDERS[what](d)
    log_action(request, "export", "Exportação %(w)s.%(f)s (%(n)s linhas)", level=30 if what == "pacientes" else 20,
               w=what, f=fmt, n=len(rows))
    name = f"{what}-{d.slug}-{timezone.localdate():%Y%m%d}"
    return _csv(name, header, rows) if fmt == "csv" else _xlsx(name, header, rows)


@dentist_required
@require_POST
def export_backup(request):
    from apps.dataio.exporter import export_bundle

    d = request.dentist
    data = export_bundle(fmt="zip", dentists=[d], include_files=True, include_passwords=False, include_platform=False)
    log_action(request, "export", "Backup da clínica %(s)s baixado", level=30, s=d.slug)
    resp = HttpResponse(data, content_type="application/zip")
    resp["Content-Disposition"] = f'attachment; filename="backup-{d.slug}-{timezone.localdate():%Y%m%d}.zip"'
    return resp
